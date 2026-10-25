import csv
import hashlib
import hmac
import io
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.evidence import append_event
from app.models import (
    DataQualityFinding,
    DataRecordRef,
    DatasetImport,
    Document,
    DocumentExtraction,
    IdentityToken,
    ImportFieldProfile,
    MatchCandidate,
    Source,
    Workspace,
)
from app.source_store import SourceStore

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_STRUCTURED_ROWS = 10_000
EMAIL_PATTERN = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
PHONE_PATTERN = re.compile(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}(?!\d)")


class IngestionProblem(Exception):
    def __init__(self, code: str, message: str, status_code: int = 422, details: dict | None = None) -> None:
        self.code = code
        self.message = message
        self.status_code = status_code
        self.details = details or {}


@dataclass
class ParsedPayload:
    kind: str
    rows: list[dict]
    headers: list[str]
    text: str | None = None


@dataclass
class Profile:
    source_field: str
    normalized_field: str
    primitive_type: str
    classification: str
    confidence: float
    detection_reason: str
    non_null_count: int
    distinct_count: int
    masked_samples: list[str]
    mapping_status: str


def normalize_field(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.strip().lower()).strip("_")


def infer_content_type(filename: str) -> str:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return {"csv": "text/csv", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "txt": "text/plain", "eml": "message/rfc822"}.get(suffix, "application/octet-stream")


def parse_payload(filename: str, payload: bytes) -> ParsedPayload:
    if not filename or "/" in filename or "\\" in filename:
        raise IngestionProblem("INVALID_FILENAME", "The uploaded filename is invalid.")
    if not payload:
        raise IngestionProblem("EMPTY_FILE", "The uploaded file is empty.")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise IngestionProblem("FILE_TOO_LARGE", "The uploaded file exceeds the 10 MiB limit.", 413)
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        return parse_csv(payload)
    if suffix == "xlsx":
        return parse_xlsx(payload)
    if suffix in {"txt", "eml"}:
        return parse_text(payload)
    raise IngestionProblem("UNSUPPORTED_FILE_TYPE", "Only CSV, XLSX, TXT, and EML imports are supported.")


def parse_csv(payload: bytes) -> ParsedPayload:
    try:
        content = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IngestionProblem("UNSUPPORTED_ENCODING", "CSV imports must use UTF-8 encoding.") from exc
    reader = csv.DictReader(io.StringIO(content), restkey="__extra__")
    headers = reader.fieldnames or []
    validate_headers(headers)
    rows: list[dict] = []
    for row in reader:
        if row.get("__extra__") is not None:
            raise IngestionProblem("INCONSISTENT_CSV_ROW", "The CSV contains rows with more values than headers.")
        rows.append({header: normalize_value(row.get(header)) for header in headers})
        if len(rows) > MAX_STRUCTURED_ROWS:
            raise IngestionProblem("ROW_LIMIT_EXCEEDED", "The import exceeds the 10,000-row limit.", 413)
    return ParsedPayload(kind="tabular", rows=rows, headers=headers)


def parse_xlsx(payload: bytes) -> ParsedPayload:
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise IngestionProblem("INVALID_XLSX", "The XLSX file could not be opened.") from exc
    if workbook.security.lockStructure:
        raise IngestionProblem("PASSWORD_PROTECTED_XLSX", "Password-protected spreadsheets are not supported.")
    sheet = next((candidate for candidate in workbook.worksheets if candidate.sheet_state == "visible"), None)
    if sheet is None:
        raise IngestionProblem("EMPTY_XLSX", "The spreadsheet has no visible worksheet.")
    iterator = sheet.iter_rows(values_only=True)
    header_row = next(iterator, None)
    headers = [str(value).strip() if value is not None else "" for value in (header_row or [])]
    validate_headers(headers)
    rows: list[dict] = []
    for values in iterator:
        if all(value is None for value in values):
            continue
        if len(values) > len(headers) and any(value is not None for value in values[len(headers):]):
            raise IngestionProblem("INCONSISTENT_XLSX_ROW", "The spreadsheet contains values beyond the header range.")
        rows.append({header: normalize_value(values[index] if index < len(values) else None) for index, header in enumerate(headers)})
        if len(rows) > MAX_STRUCTURED_ROWS:
            raise IngestionProblem("ROW_LIMIT_EXCEEDED", "The import exceeds the 10,000-row limit.", 413)
    return ParsedPayload(kind="tabular", rows=rows, headers=headers)


def parse_text(payload: bytes) -> ParsedPayload:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IngestionProblem("UNSUPPORTED_ENCODING", "Text imports must use UTF-8 encoding.") from exc
    if not text.strip():
        raise IngestionProblem("EMPTY_FILE", "The uploaded file is empty.")
    return ParsedPayload(kind="text", rows=[{"content": text}], headers=["content"], text=text)


def validate_headers(headers: list[str]) -> None:
    if not headers or any(not header.strip() for header in headers):
        raise IngestionProblem("INVALID_HEADERS", "Structured imports require a non-empty header row.")
    normalized = [normalize_field(header) for header in headers]
    if len(set(normalized)) != len(normalized):
        raise IngestionProblem("DUPLICATE_HEADERS", "Structured imports cannot contain duplicate headers.")


def normalize_value(value: object | None) -> str | None:
    if value is None:
        return None
    return str(value).strip() or None


def hmac_token(token_type: str, value: str) -> str:
    return hmac.new(get_settings().hmac_secret.encode(), f"{token_type}:{value}".encode(), hashlib.sha256).hexdigest()


def normalized_email(value: str) -> str | None:
    value = value.strip().lower()
    return value if EMAIL_PATTERN.fullmatch(value) else None


def normalized_phone(value: str) -> str | None:
    digits = re.sub(r"\D", "", value)
    if len(digits) == 10 and digits[0] in "6789":
        return f"+91{digits}"
    if len(digits) == 12 and digits.startswith("91") and digits[2] in "6789":
        return f"+{digits}"
    return None


def mask_value(value: str, classification: str) -> str:
    if classification == "DIRECT_IDENTIFIER":
        email = normalized_email(value)
        if email:
            local, domain = email.split("@", 1)
            return f"{local[:1]}***@{domain}"
        phone = normalized_phone(value)
        if phone:
            return f"******{phone[-4:]}"
        return "[redacted identifier]"
    if classification in {"PERSON_NAME", "ADDRESS", "DATE_OF_BIRTH"}:
        return "[redacted personal field]"
    if classification == "COMMUNICATION_CONTENT":
        return "[redacted content]"
    return "[available in source]"


def profile_fields(payload: ParsedPayload) -> list[Profile]:
    profiles: list[Profile] = []
    for header in payload.headers:
        values = [row.get(header) for row in payload.rows if row.get(header) is not None]
        classification, confidence, reason, mapping_status = classify_field(header, [value for value in values if value])
        profiles.append(Profile(
            source_field=header,
            normalized_field=normalize_field(header),
            primitive_type=infer_primitive_type([value for value in values if value]),
            classification=classification,
            confidence=confidence,
            detection_reason=reason,
            non_null_count=len(values),
            distinct_count=len(set(values)),
            masked_samples=[mask_value(value, classification) for value in values[:3]],
            mapping_status=mapping_status,
        ))
    return profiles


def classify_field(header: str, values: list[str]) -> tuple[str, float, str, str]:
    normalized = normalize_field(header)
    if normalized in {"email", "email_address", "mobile", "phone", "phone_number", "contact_number"}:
        return "DIRECT_IDENTIFIER", 1.0, "recognized header", "MAPPED"
    if normalized in {"subject_alias", "customer_id", "learner_id", "record_id"}:
        return "PSEUDONYMOUS_IDENTIFIER", 0.95, "recognized identifier header", "MAPPED"
    if normalized in {"name", "full_name", "first_name", "last_name", "guardian_name"} or normalized.endswith("_name"):
        return "PERSON_NAME", 0.95, "recognized name header", "MAPPED"
    if normalized in {"address", "postal_address", "city"}:
        return "ADDRESS", 0.9, "recognized address header", "MAPPED"
    if normalized in {"dob", "date_of_birth", "birth_date"}:
        return "DATE_OF_BIRTH", 0.95, "recognized birth-date header", "MAPPED"
    if normalized in {"message", "body", "content", "notes", "note", "comments"}:
        return "COMMUNICATION_CONTENT", 0.95, "recognized free-text header", "MAPPED"
    if any(token in normalized for token in ("consent", "authorization", "authorisation", "opt_in")):
        return "AUTHORIZATION_EVIDENCE", 0.9, "recognized authorization header", "MAPPED"
    if "purpose" in normalized:
        return "PURPOSE_METADATA", 0.9, "recognized purpose header", "MAPPED"
    if normalized == "guardian_contact":
        return "DIRECT_IDENTIFIER", 0.65, "possible contact field requires confirmation", "REVIEW_REQUIRED"
    if values and sum(bool(normalized_email(value) or normalized_phone(value)) for value in values) / len(values) >= 0.8:
        return "DIRECT_IDENTIFIER", 0.85, "value pattern suggests direct identifiers", "REVIEW_REQUIRED"
    return "UNKNOWN", 0.0, "no deterministic mapping rule matched", "REVIEW_REQUIRED"


def infer_primitive_type(values: list[str]) -> str:
    if not values:
        return "string"
    lowered = {value.lower() for value in values}
    if lowered <= {"true", "false", "yes", "no"}:
        return "boolean"
    try:
        [int(value) for value in values]
        return "integer"
    except ValueError:
        pass
    try:
        [Decimal(value) for value in values]
        return "decimal"
    except (InvalidOperation, ValueError):
        pass
    try:
        [datetime.fromisoformat(value.replace("Z", "+00:00")) for value in values]
        return "datetime"
    except ValueError:
        return "string"


def redacted_text_preview(text: str) -> tuple[str, dict[str, int]]:
    emails = EMAIL_PATTERN.findall(text)
    phones = PHONE_PATTERN.findall(text)
    redacted = EMAIL_PATTERN.sub("[redacted email]", text)
    redacted = PHONE_PATTERN.sub("[redacted phone]", redacted)
    return redacted[:2_000], {"emails": len(emails), "phones": len(phones)}


def find_or_create_document(db: Session, workspace: Workspace, source: Source, filename: str, digest: str, kind: str) -> Document:
    document = db.scalar(
        select(Document)
        .where(Document.workspace_id == workspace.id, Document.source_id == source.id, Document.filename == filename)
        .order_by(Document.created_at.desc())
    )
    if document and db.scalar(select(DatasetImport.id).where(DatasetImport.document_id == document.id)) is None:
        return document
    return Document(
        workspace_id=workspace.id,
        source_id=source.id,
        document_key=f"import-{digest[:16]}",
        synthetic_label="synthetic_demo_data",
        document_type="structured_dataset" if kind == "tabular" else "text_document",
        filename=filename,
        mime_type=infer_content_type(filename),
        source_path=f"source-store://{source.id}/{digest}",
        display_title=filename,
        document_created_at=datetime.now(timezone.utc),
        processing_status="processing",
        metadata_json={"synthetic": True, "uploaded": True},
    )


def direct_tokens(row: dict, profiles: list[Profile]) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    profile_by_field = {profile.source_field: profile for profile in profiles}
    for field, value in row.items():
        if not value:
            continue
        profile = profile_by_field[field]
        if profile.classification == "DIRECT_IDENTIFIER":
            if email := normalized_email(value):
                tokens.append(("email", hmac_token("email", email)))
            elif phone := normalized_phone(value):
                tokens.append(("phone", hmac_token("phone", phone)))
        elif profile.classification == "PSEUDONYMOUS_IDENTIFIER":
            tokens.append(("pseudonymous_identifier", hmac_token("pseudonymous_identifier", value.lower())))
    return sorted(set(tokens))


def name_tokens(row: dict, profiles: list[Profile]) -> list[tuple[str, str]]:
    profile_by_field = {profile.source_field: profile for profile in profiles}
    return [("name", hmac_token("name", value.lower().strip())) for field, value in row.items() if value and profile_by_field[field].classification == "PERSON_NAME"]


def _existing_subject_refs(db: Session, workspace_id: str, tokens: list[tuple[str, str]]) -> set[str]:
    record_ids: set[str] = set()
    for token_type, token in tokens:
        record_ids.update(db.scalars(select(IdentityToken.record_ref_id).where(IdentityToken.workspace_id == workspace_id, IdentityToken.token_type == token_type, IdentityToken.hmac_token == token)))
    if not record_ids:
        return set()
    return {value for value in db.scalars(select(DataRecordRef.subject_ref).where(DataRecordRef.id.in_(record_ids), DataRecordRef.subject_ref.is_not(None))) if value}


def create_match_candidates(db: Session, workspace: Workspace, record: DataRecordRef, weak_tokens: list[tuple[str, str]]) -> int:
    created = 0
    for token_type, token in weak_tokens:
        prior_ids = db.scalars(select(IdentityToken.record_ref_id).where(IdentityToken.workspace_id == workspace.id, IdentityToken.token_type == token_type, IdentityToken.hmac_token == token)).all()
        for prior_id in prior_ids:
            if prior_id == record.id:
                continue
            left, right = sorted((prior_id, record.id))
            exists = db.scalar(select(MatchCandidate.id).where(MatchCandidate.left_record_ref_id == left, MatchCandidate.right_record_ref_id == right))
            if exists is None:
                db.add(MatchCandidate(workspace_id=workspace.id, left_record_ref_id=left, right_record_ref_id=right, match_method="exact_name_token", confidence=0.55, status="REVIEW_REQUIRED"))
                created += 1
    return created


def add_quality_findings(db: Session, workspace: Workspace, source: Source, import_row: DatasetImport, profiles: list[Profile], candidate_count: int) -> None:
    def add(code: str, severity: str, count: int, field_name: str | None = None) -> None:
        db.add(DataQualityFinding(workspace_id=workspace.id, source_id=source.id, dataset_import_id=import_row.id, code=code, severity=severity, field_name=field_name, affected_record_count=count, status="OPEN", details_json={}))
    for profile in profiles:
        if profile.mapping_status == "REVIEW_REQUIRED":
            add("UNCERTAIN_FIELD_MAPPING", "warning", import_row.row_count or 0, profile.source_field)
    purpose_profiles = [profile for profile in profiles if profile.classification == "PURPOSE_METADATA"]
    authorization_profiles = [profile for profile in profiles if profile.classification == "AUTHORIZATION_EVIDENCE"]
    contact_profiles = [profile for profile in profiles if profile.classification == "DIRECT_IDENTIFIER"]
    if not purpose_profiles:
        add("MISSING_PURPOSE_METADATA", "warning", import_row.row_count or 0)
    elif missing := max((import_row.row_count or 0) - max(profile.non_null_count for profile in purpose_profiles), 0):
        add("MISSING_PURPOSE_METADATA", "warning", missing)
    if not authorization_profiles:
        add("MISSING_AUTHORIZATION_EVIDENCE", "info", import_row.row_count or 0)
    if not contact_profiles:
        add("INCOMPLETE_CONTACT_DATA", "warning", import_row.row_count or 0)
    if candidate_count:
        add("AMBIGUOUS_MATCH_CANDIDATE", "warning", candidate_count)


def ingest_file(db: Session, workspace: Workspace, source: Source, filename: str, payload: bytes) -> tuple[DatasetImport, bool]:
    digest = hashlib.sha256(payload).hexdigest()
    existing = db.scalar(select(DatasetImport).where(DatasetImport.source_id == source.id, DatasetImport.content_sha256 == digest))
    if existing is not None:
        return existing, True

    import_row: DatasetImport | None = None
    store: SourceStore | None = None
    try:
        suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        provisional_kind = "tabular" if suffix in {"csv", "xlsx"} else "text"
        document = find_or_create_document(db, workspace, source, filename, digest, provisional_kind)
        db.add(document)
        db.flush()
        import_row = DatasetImport(workspace_id=workspace.id, source_id=source.id, document_id=document.id, filename=filename, content_type=infer_content_type(filename), byte_size=len(payload), content_sha256=digest, status="RECEIVED")
        db.add(import_row)
        db.flush()
        parsed = parse_payload(filename, payload)
        import_row.status = "VALIDATING"
        profiles = profile_fields(parsed)
        import_row.status = "PROCESSING"
        store = SourceStore(source.id)
        records = [(f"{import_row.id}-row-{index + 1}", index + 1, row) for index, row in enumerate(parsed.rows)]
        store.write_records(import_row.id, records)
        if parsed.text:
            store.write_document(document.id, import_row.id, parsed.text)

        for ordinal, profile in enumerate(profiles, start=1):
            db.add(ImportFieldProfile(dataset_import_id=import_row.id, ordinal=ordinal, source_field=profile.source_field, normalized_field=profile.normalized_field, primitive_type=profile.primitive_type, classification=profile.classification, confidence=profile.confidence, detection_reason=profile.detection_reason, non_null_count=profile.non_null_count, distinct_count=profile.distinct_count, masked_samples_json=profile.masked_samples, mapping_status=profile.mapping_status))

        candidate_count = 0
        for source_record_id, _, row in records:
            direct = direct_tokens(row, profiles)
            known_refs = _existing_subject_refs(db, workspace.id, direct)
            if len(known_refs) > 1:
                subject_ref, match_status = None, "CONFLICT"
            elif known_refs:
                subject_ref, match_status = next(iter(known_refs)), "MATCHED"
            elif direct:
                subject_ref, match_status = f"subref_{direct[0][1][:20]}", "MATCHED"
            else:
                subject_ref, match_status = None, "UNMATCHED"
            field_metadata = {profile.normalized_field: {"classification": profile.classification, "present": row.get(profile.source_field) is not None} for profile in profiles}
            record = DataRecordRef(workspace_id=workspace.id, source_id=source.id, dataset_import_id=import_row.id, source_record_id=source_record_id, subject_ref=subject_ref, record_type="tabular_row" if parsed.kind == "tabular" else "text_document", match_status=match_status, field_metadata_json=field_metadata)
            db.add(record)
            db.flush()
            weak = name_tokens(row, profiles)
            for token_type, token in direct + weak:
                db.add(IdentityToken(workspace_id=workspace.id, record_ref_id=record.id, token_type=token_type, hmac_token=token))
            db.flush()
            candidate_count += create_match_candidates(db, workspace, record, weak)

        if parsed.text:
            preview, entities = redacted_text_preview(parsed.text)
            db.add(DocumentExtraction(document_id=document.id, dataset_import_id=import_row.id, extraction_status="COMPLETE", character_count=len(parsed.text), redacted_text_preview=preview, entity_summary_json=entities, warnings_json=[]))
        import_row.status = "COMPLETE"
        import_row.row_count = len(records)
        import_row.completed_at = datetime.now(timezone.utc)
        document.processing_status = "catalogued"
        source.last_synced_at = datetime.now(timezone.utc)
        add_quality_findings(db, workspace, source, import_row, profiles, candidate_count)
        append_event(
            db,
            workspace.id,
            "IMPORT_COMPLETED",
            {
                "source_id": source.id,
                "source_key": source.source_key,
                "import_id": import_row.id,
                "row_count": len(records),
                "document_type": document.document_type,
            },
            entity_type="dataset_import",
            entity_ref=import_row.id,
            occurred_at=import_row.completed_at,
        )
        db.commit()
        db.refresh(import_row)
        return import_row, False
    except IngestionProblem:
        if import_row is not None:
            import_row.status = "FAILED"
            import_row.error_code = "IMPORT_VALIDATION_FAILED"
            import_row.error_message = "The file did not pass import validation."
            db.commit()
        raise
    except Exception as exc:
        db.rollback()
        if store is not None and import_row is not None:
            store.delete_import(import_row.id)
        raise IngestionProblem("IMPORT_PROCESSING_FAILED", "The import could not be completed.", 500) from exc
